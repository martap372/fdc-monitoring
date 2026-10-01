import json
import os

import pandas as pd
from openpyxl import Workbook, load_workbook

from pt import BACKEND_DIR, CLIENT_SINCE_PATH, WORKBOOK_LOCK


MONTHS = [
	'Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho',
	'Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro'
]
HEADERS = ['Nº de Sócio', 'Nome do Cliente', 'Nº Avaliação', 'Avaliação Física', 'Plano de Treino', 'Dias de Diferença', 'Cliente PT', 'Fecho']


def af_filename(month, year):
	return f'AF_{MONTHS[month - 1]}{year}.xlsx'


def af_files():
	return sorted(
		filename for filename in os.listdir(BACKEND_DIR)
		if filename.startswith('AF_') and filename.endswith('.xlsx')
	)


def _client_since_data():
	with open(CLIENT_SINCE_PATH, encoding='utf-8') as client_since_file:
		return {
			_member_key(member_number): starting_month
			for member_number, starting_month in json.load(client_since_file).items()
		}


def _af_flags(member_number, evaluation_date, client_since):
	starting_month = client_since.get(_member_key(member_number))
	client_pt = 'Sim' if starting_month is not None else 'Não'
	parsed_date = pd.to_datetime(evaluation_date, dayfirst=True, errors='coerce')
	parsed_start = pd.to_datetime(starting_month, format='%m-%Y', errors='coerce')
	fecho = (
		'Sim'
		if starting_month is not None
		and not pd.isna(parsed_date)
		and not pd.isna(parsed_start)
		and parsed_date.month == parsed_start.month
		and parsed_date.year == parsed_start.year
		else 'Não'
	)
	return client_pt, fecho


def regenerate_af_flags():
	client_since = _client_since_data()
	for filename in af_files():
		workbook_path = os.path.join(BACKEND_DIR, filename)
		workbook = load_workbook(workbook_path)
		changed = False
		for worksheet in workbook.worksheets:
			headers = [cell.value for cell in worksheet[1]]
			try:
				member_column = headers.index('Nº de Sócio') + 1
				date_column = headers.index('Avaliação Física') + 1
			except ValueError:
				continue
			client_pt_column = headers.index('Cliente PT') + 1 if 'Cliente PT' in headers else worksheet.max_column + 1
			fecho_column = headers.index('Fecho') + 1 if 'Fecho' in headers else worksheet.max_column + (1 if client_pt_column <= worksheet.max_column else 2)
			if worksheet.cell(1, client_pt_column).value != 'Cliente PT':
				worksheet.cell(1, client_pt_column, 'Cliente PT')
				changed = True
			if worksheet.cell(1, fecho_column).value != 'Fecho':
				worksheet.cell(1, fecho_column, 'Fecho')
				changed = True
			for row in range(2, worksheet.max_row + 1):
				flags = _af_flags(
					worksheet.cell(row, member_column).value,
					worksheet.cell(row, date_column).value,
					client_since,
				)
				for column, value in zip((client_pt_column, fecho_column), flags):
					cell = worksheet.cell(row, column)
					if cell.value != value:
						cell.value = value
						changed = True
			worksheet.auto_filter.ref = worksheet.dimensions
		if changed:
			temporary_path = f'{workbook_path}.tmp'
			workbook.save(temporary_path)
			os.replace(temporary_path, workbook_path)
		workbook.close()


def _json_value(value):
	if value is None or pd.isna(value):
		return None
	if hasattr(value, 'isoformat'):
		return value.isoformat().split('T')[0]
	return value.item() if hasattr(value, 'item') else value


def af_workbook_data(filename):
	workbook_path = os.path.join(BACKEND_DIR, filename)
	workbook = load_workbook(workbook_path, read_only=True, data_only=False)
	result = {}
	for worksheet in workbook.worksheets:
		result[worksheet.title] = [
			[_json_value(value) for value in row]
			for row in worksheet.iter_rows(values_only=True)
		]
	workbook.close()
	return result


def _clean_value(value):
	value = _json_value(value)
	if isinstance(value, float) and value.is_integer():
		return int(value)
	return value


def _member_key(value):
	value = _clean_value(value)
	if value is None or value == '':
		return None
	return str(value).strip()


def _existing_training_plans(filename, person):
	workbook_path = os.path.join(BACKEND_DIR, filename)
	if not os.path.exists(workbook_path):
		return {}

	workbook = load_workbook(workbook_path, read_only=True, data_only=False)
	if person not in workbook.sheetnames:
		workbook.close()
		return {}

	worksheet = workbook[person]
	plans = {
		_member_key(row[0]): row[4]
		for row in worksheet.iter_rows(min_row=2, values_only=True)
		if len(row) > 4 and _member_key(row[0]) is not None
	}
	workbook.close()
	return plans


def _mark_af_plans(planos_filenames, af_filename, person):
	if isinstance(planos_filenames, str):
		planos_filenames = [planos_filenames]
	training_dates = {}
	for planos_filename in planos_filenames:
		planos_path = os.path.join(BACKEND_DIR, planos_filename)
		if not os.path.exists(planos_path):
			continue
		planos_workbook = load_workbook(planos_path, read_only=True, data_only=False)
		if person in planos_workbook.sheetnames:
			planos_worksheet = planos_workbook[person]
			headers = [cell.value for cell in planos_worksheet[1]]
			try:
				member_column = headers.index('Nº de Sócio')
				date_column = headers.index('Data') + 1
			except ValueError:
				member_column = date_column = None
			if member_column is not None:
				for row in planos_worksheet.iter_rows(min_row=2, values_only=True):
					member_number = _member_key(row[member_column])
					plan_date = pd.to_datetime(row[date_column - 1], dayfirst=True, errors='coerce')
					if member_number is not None and not pd.isna(plan_date):
						training_dates.setdefault(member_number, []).append(plan_date)
		planos_workbook.close()

	af_path = os.path.join(BACKEND_DIR, af_filename)
	workbook = load_workbook(af_path)
	if person not in workbook.sheetnames:
		workbook.close()
		return

	worksheet = workbook[person]
	headers = [cell.value for cell in worksheet[1]]
	try:
		member_column = headers.index('Nº de Sócio') + 1
		date_column = headers.index('Avaliação Física') + 1 if 'Avaliação Física' in headers else headers.index('Data') + 1
		plan_column = headers.index('Plano de Treino') + 1
		difference_column = plan_column + 1
	except ValueError:
		workbook.close()
		return

	if worksheet.cell(1, difference_column).value != 'Dias de Diferença':
		worksheet.cell(1, difference_column, 'Dias de Diferença')

	changed = False
	for row in range(2, worksheet.max_row + 1):
		member_number = _member_key(worksheet.cell(row, member_column).value)
		af_date = pd.to_datetime(worksheet.cell(row, date_column).value, dayfirst=True, errors='coerce')
		worksheet.cell(row, date_column).number_format = 'dd-mm-yyyy'
		cell = worksheet.cell(row, plan_column)
		difference_cell = worksheet.cell(row, difference_column)
		matching_dates = []
		if member_number is not None and not pd.isna(af_date):
			matching_dates = [
				plan_date for plan_date in training_dates.get(member_number, [])
				if plan_date >= af_date - pd.Timedelta(days=1)
			]
		value = min(matching_dates).strftime('%d-%m-%Y') if matching_dates else '-'
		if cell.value != value:
			cell.value = value
			if value != '-':
				cell.number_format = 'dd-mm-yyyy'
			changed = True
		if matching_dates:
			plan_value = min(matching_dates)
			difference_value = int((plan_value - af_date).days)
		else:
			difference_value = '-'
		if difference_cell.value != difference_value:
			difference_cell.value = difference_value
			changed = True

	if changed:
		temporary_path = f'{af_path}.tmp'
		workbook.save(temporary_path)
		os.replace(temporary_path, af_path)
	workbook.close()


def _import_period(values):
	periods = set()
	for value in values:
		if value is None or (isinstance(value, str) and not value.strip()):
			continue
		parsed = pd.to_datetime(value, dayfirst=True, errors='coerce')
		if pd.isna(parsed):
			continue
		periods.add((parsed.month, parsed.year))

	if not periods:
		raise ValueError('Não foi possível identificar o mês e o ano na coluna Inserida em.')
	if len(periods) > 1:
		raise ValueError('A coluna Inserida em contém mais do que um mês ou ano.')
	return periods.pop()


def import_(filename, person):
	import app

	if person not in app.PEOPLE:
		raise ValueError('PT inválido.')

	input_path = os.path.join(BACKEND_DIR, filename)
	source = pd.read_excel(input_path)
	source.columns = [str(column).strip() for column in source.columns]
	required_columns = ['Código', 'Cliente', '#', 'Inserida em']
	missing_columns = [column for column in required_columns if column not in source.columns]
	if missing_columns:
		raise ValueError(f'Colunas em falta: {", ".join(missing_columns)}.')
	month, year = _import_period(source['Inserida em'])
	client_since = _client_since_data()
	valid_rows = source[
		source['Código'].notna()
		& source['Inserida em'].map(lambda value: not pd.isna(pd.to_datetime(value, dayfirst=True, errors='coerce')))
	]
	if valid_rows.empty:
		raise ValueError('O ficheiro AF não contém registos válidos.')

	rows_by_member = {}
	for _, source_row in valid_rows.iterrows():
		member_number = _clean_value(source_row['Código'])
		member_key = _member_key(member_number)
		evaluation_date = _clean_value(source_row['Inserida em'])
		parsed_date = pd.to_datetime(evaluation_date, dayfirst=True, errors='coerce')
		client_pt, fecho = _af_flags(member_number, evaluation_date, client_since)
		row = [
			member_number,
			_clean_value(source_row['Cliente']),
			_clean_value(source_row['#']),
			evaluation_date,
			'-',
			'-',
			client_pt,
			fecho,
		]
		current = rows_by_member.get(member_key)
		if current is None or parsed_date < current[0]:
			rows_by_member[member_key] = (parsed_date, row)
	rows = [row for _, row in rows_by_member.values()]

	output_filename = af_filename(month, year)
	output_path = os.path.join(BACKEND_DIR, output_filename)
	with WORKBOOK_LOCK:
		if os.path.exists(output_path):
			workbook = load_workbook(output_path)
		else:
			workbook = Workbook()
			workbook.remove(workbook.active)

		if person in workbook.sheetnames:
			del workbook[person]
		worksheet = workbook.create_sheet(person)
		worksheet.append(HEADERS)
		for row in rows:
			worksheet.append(row)
		worksheet.freeze_panes = 'A2'
		worksheet.auto_filter.ref = worksheet.dimensions
		for column in worksheet.columns:
			width = max(len(str(cell.value or '')) for cell in column) + 2
			worksheet.column_dimensions[column[0].column_letter].width = min(width, 35)

		temporary_path = f'{output_path}.tmp'
		workbook.save(temporary_path)
		os.replace(temporary_path, output_path)
		workbook.close()
		planos_filenames = [f'Planos_{MONTHS[month - 1]}{year}.xlsx']
		if month < len(MONTHS):
			planos_filenames.append(f'Planos_{MONTHS[month]}{year}.xlsx')
		else:
			planos_filenames.append(f'Planos_{MONTHS[0]}{year + 1}.xlsx')
		_mark_af_plans(planos_filenames, output_filename, person)

	return output_filename
