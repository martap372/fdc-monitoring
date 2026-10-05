import json
import ipaddress
import os
import re
import subprocess
from copy import copy
from urllib.parse import urlparse

from flask import Flask, jsonify, request, send_file
from werkzeug.utils import secure_filename

PEOPLE = (
    'André Mota',
    'Daniel Araújo',
    'Emanuel Ferreira',
    'Pedro Freitas',
    'Rúben Ramos',
    'Simão Sá',
)

from mapa_sala import (
    WEEKLY_HOURS_FILE,
    add_day_of_week,
    build_monthly_hours,
    read_events,
    schedule_vencimento_update,
    vencimento_update_status,
    write_events,
)
from pt import BACKEND_DIR, WORKBOOK_LOCK, import_, pt_client_status, pt_files, workbook_data
from af import MONTHS as AF_MONTHS, af_files, af_workbook_data, import_ as import_af
from planos import planos_files, planos_workbook_data, import_ as import_planos
from convert import convert_xls
from vencimento import (
    generate_vencimento,
    get_af_closure_override,
    set_af_closure_override,
)
from openpyxl import load_workbook
from uuid import uuid4


app = Flask(__name__)


def vencimento_files():
    for filename in os.listdir(BACKEND_DIR):
        if (
            filename.startswith('Vencimento_')
            and filename.endswith('.xlsx')
            and '.tmp.' in filename
        ):
            tmp_path = os.path.join(BACKEND_DIR, filename)
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    files = sorted(
        filename for filename in os.listdir(BACKEND_DIR)
        if filename.startswith('Vencimento_')
        and filename.endswith('.xlsx')
        and '.tmp.' not in filename
    )
    if not files:
        return [generate_vencimento()]
    return files


@app.after_request
def add_cors_headers(response):
    origin = request.headers.get('Origin')
    parsed_origin = urlparse(origin) if origin else None
    hostname = parsed_origin.hostname if parsed_origin else None
    is_localhost = hostname in {'localhost', '127.0.0.1', '::1', '0.0.0.0'}
    is_local_hostname = hostname and ('.' not in hostname or hostname.endswith('.local'))
    try:
        is_private_ip = parsed_origin and ipaddress.ip_address(parsed_origin.hostname).is_private
    except ValueError:
        is_private_ip = False
    if origin and parsed_origin.scheme in {'http', 'https'} and (is_localhost or is_private_ip or is_local_hostname):
        response.headers['Access-Control-Allow-Origin'] = origin
        response.headers['Access-Control-Allow-Private-Network'] = 'true'
    response.headers['Access-Control-Allow-Headers'] = 'Content-Type'
    response.headers['Access-Control-Allow-Methods'] = 'GET, POST, PUT, DELETE, OPTIONS'
    return response


@app.get('/api/pt-files')
def get_pt_files():
    return jsonify(pt_files())


@app.get('/api/pt-files/<filename>')
def get_pt_workbook(filename):
    if filename not in pt_files():
        return jsonify({'error': 'Ficheiro PT não encontrado'}), 404
    return jsonify(workbook_data(filename))


@app.get('/api/pt-files/<filename>/download')
def download_pt_workbook(filename):
    if filename not in pt_files():
        return jsonify({'error': 'Ficheiro PT não encontrado'}), 404
    return send_file(os.path.join(BACKEND_DIR, filename), as_attachment=True, download_name=filename)


@app.get('/api/af-files')
def get_af_files():
    return jsonify(af_files())


@app.get('/api/af-files/<filename>')
def get_af_workbook(filename):
    if filename not in af_files():
        return jsonify({'error': 'Ficheiro AF não encontrado'}), 404
    return jsonify(af_workbook_data(filename))


@app.get('/api/af-files/<filename>/download')
def download_af_workbook(filename):
    if filename not in af_files():
        return jsonify({'error': 'Ficheiro AF não encontrado'}), 404
    return send_file(os.path.join(BACKEND_DIR, filename), as_attachment=True, download_name=filename)


@app.route('/api/af-files/<filename>/fechos', methods=['GET', 'PUT'])
def update_af_fechos(filename):
    if filename not in af_files():
        return jsonify({'error': 'Ficheiro AF não encontrado'}), 404
    match = re.fullmatch(r'AF_(.+)(\d{4})\.xlsx', filename)
    if not match or match.group(1) not in AF_MONTHS:
        return jsonify({'error': 'Ficheiro AF inválido'}), 400

    changes = request.get_json(silent=True) or {} if request.method == 'PUT' else {}
    sheet_name = changes.get('sheet') if request.method == 'PUT' else request.args.get('sheet')
    if sheet_name not in PEOPLE:
        return jsonify({'error': 'Folha AF inválida'}), 400
    year = int(match.group(2))
    month = AF_MONTHS.index(match.group(1)) + 1
    if request.method == 'GET':
        return jsonify({'value': get_af_closure_override(year, month, sheet_name)})

    value = changes.get('value')
    try:
        numeric_value = float(value)
        if not numeric_value.is_integer() or numeric_value < 0:
            raise ValueError
        numeric_value = int(numeric_value)
    except (TypeError, ValueError):
        return jsonify({'error': 'Fechos deve ser um número inteiro não negativo'}), 400

    with WORKBOOK_LOCK:
        set_af_closure_override(year, month, sheet_name, numeric_value)
    schedule_vencimento_update(year=year)
    return jsonify({'value': numeric_value})


@app.get('/api/planos-files')
def get_planos_files():
    return jsonify(planos_files())


@app.get('/api/planos-files/<filename>')
def get_planos_workbook(filename):
    if filename not in planos_files():
        return jsonify({'error': 'Ficheiro de planos não encontrado'}), 404
    return jsonify(planos_workbook_data(filename))


@app.get('/api/planos-files/<filename>/download')
def download_planos_workbook(filename):
    if filename not in planos_files():
        return jsonify({'error': 'Ficheiro de planos não encontrado'}), 404
    return send_file(os.path.join(BACKEND_DIR, filename), as_attachment=True, download_name=filename)


@app.post('/api/af-import')
def import_af_workbook():
    uploaded_file = request.files.get('file')
    if not uploaded_file or not uploaded_file.filename.lower().endswith(('.xlsx', '.xls')):
        return jsonify({'error': 'Envie um ficheiro Excel (.xlsx ou .xls)'}), 400

    person = request.form.get('person', '')
    filename = secure_filename(uploaded_file.filename)
    input_path = os.path.join(BACKEND_DIR, filename)
    converted_filename = None
    uploaded_file.save(input_path)
    try:
        import_filename, converted = convert_xls(input_path)
        converted_filename = os.path.basename(import_filename) if converted else None
        output_filename = import_af(import_filename, person)
        year_match = re.fullmatch(r'AF_.+(\d{4})\.xlsx', output_filename)
        if year_match:
            generate_vencimento(year=int(year_match.group(1)))
    except (ValueError, KeyError, IndexError, OSError, RuntimeError, subprocess.TimeoutExpired) as import_error:
        return jsonify({'error': str(import_error)}), 400
    finally:
        if os.path.exists(input_path):
            os.remove(input_path)
        if converted_filename:
            converted_path = os.path.join(BACKEND_DIR, converted_filename)
            if os.path.exists(converted_path):
                os.remove(converted_path)

    return jsonify({'filename': output_filename}), 201


@app.post('/api/planos-import')
def import_planos_workbook():
    uploaded_file = request.files.get('file')
    if not uploaded_file or not uploaded_file.filename.lower().endswith(('.xlsx', '.xls')):
        return jsonify({'error': 'Envie um ficheiro Excel (.xlsx ou .xls)'}), 400
    person = request.form.get('person', '')
    filename = secure_filename(uploaded_file.filename)
    input_path = os.path.join(BACKEND_DIR, filename)
    converted_filename = None
    uploaded_file.save(input_path)
    try:
        import_filename, converted = convert_xls(input_path)
        converted_filename = os.path.basename(import_filename) if converted else None
        output_filename = import_planos(os.path.basename(import_filename), person)
    except (ValueError, KeyError, IndexError, OSError, RuntimeError, subprocess.TimeoutExpired) as import_error:
        return jsonify({'error': str(import_error)}), 400
    finally:
        if os.path.exists(input_path):
            os.remove(input_path)
        if converted_filename:
            converted_path = os.path.join(BACKEND_DIR, converted_filename)
            if os.path.exists(converted_path):
                os.remove(converted_path)
    return jsonify({'filename': output_filename}), 201


@app.put('/api/af-files/<filename>/cells')
def update_af_cell(filename):
    if filename not in af_files():
        return jsonify({'error': 'Ficheiro AF não encontrado'}), 404

    changes = request.get_json(silent=True) or {}
    sheet_name = changes.get('sheet')
    row = changes.get('row')
    column = changes.get('column')
    value = changes.get('value')
    if not isinstance(sheet_name, str) or not isinstance(row, int) or column != 4:
        return jsonify({'error': 'Dados da célula inválidos'}), 400
    if row < 1 or not isinstance(value, str) or not value.strip():
        return jsonify({'error': 'Plano de Treino inválido'}), 400

    workbook_path = os.path.join(BACKEND_DIR, filename)
    with WORKBOOK_LOCK:
        workbook = load_workbook(workbook_path)
        if sheet_name not in workbook.sheetnames:
            return jsonify({'error': 'Folha não encontrada'}), 404
        cell = workbook[sheet_name].cell(row=row + 1, column=column + 1)
        cell.value = value
        temporary_path = f'{workbook_path}.tmp'
        workbook.save(temporary_path)
        os.replace(temporary_path, workbook_path)
        workbook.close()

    year_match = re.fullmatch(r'AF_.+(\d{4})\.xlsx', filename)
    if year_match:
        generate_vencimento(year=int(year_match.group(1)))

    return jsonify({'sheet': sheet_name, 'row': row, 'column': column, 'value': value})


@app.get('/api/vencimento-files')
def get_vencimento_files():
    return jsonify(vencimento_files())


@app.get('/api/vencimento-status')
def get_vencimento_status():
    return jsonify(vencimento_update_status())


@app.get('/api/vencimento-files/<filename>')
def get_vencimento_workbook(filename):
    if filename not in vencimento_files():
        return jsonify({'error': 'Ficheiro de vencimentos não encontrado'}), 404
    return jsonify(workbook_data(filename, include_styles=True, data_only=True))


@app.get('/api/vencimento-files/<filename>/download')
def download_vencimento_workbook(filename):
    if filename not in vencimento_files():
        return jsonify({'error': 'Ficheiro de vencimentos não encontrado'}), 404
    return send_file(os.path.join(BACKEND_DIR, filename), as_attachment=True, download_name=filename)


@app.put('/api/vencimento-files/<filename>/cells')
def update_vencimento_cell(filename):
    if filename not in vencimento_files():
        return jsonify({'error': 'Ficheiro de vencimentos não encontrado'}), 404

    changes = request.get_json(silent=True) or {}
    sheet_name = changes.get('sheet')
    row = changes.get('row')
    column = changes.get('column')
    allowed_columns = (
        {12, 15}
        if sheet_name == 'Pedro Freitas'
        else {14} if sheet_name == 'Simão Sá'
        else {13}
    )
    if not isinstance(sheet_name, str) or not isinstance(row, int) or column not in allowed_columns:
        return jsonify({'error': 'Dados da célula inválidos'}), 400
    if row < 2 or row > 13:
        return jsonify({'error': 'Coordenadas da célula inválidas'}), 400
    value = changes.get('value')
    try:
        value = float(value)
        if value < 0:
            raise ValueError
        if column == 12 and not value.is_integer():
            raise ValueError
        value = int(value) if column == 12 else round(value, 2)
    except (TypeError, ValueError):
        message = (
            'Nº Aulas do RUN CLUB deve ser um número inteiro não negativo'
            if column == 12
            else 'Aulas Externas deve ser um valor não negativo'
        )
        return jsonify({'error': message}), 400

    workbook_path = os.path.join(BACKEND_DIR, filename)
    with WORKBOOK_LOCK:
        workbook = load_workbook(workbook_path, data_only=False)
        if sheet_name not in workbook.sheetnames:
            return jsonify({'error': 'Folha não encontrada'}), 404

        cell = workbook[sheet_name].cell(row=row + 1, column=column + 1)
        cell.value = value
        temporary_path = f'{workbook_path}.tmp'
        workbook.save(temporary_path)
        os.replace(temporary_path, workbook_path)

    year_match = re.fullmatch(r'Vencimento_(\d{4})\.xlsx', filename)
    if year_match:
        generate_vencimento(year=int(year_match.group(1)))

    return jsonify({'sheet': sheet_name, 'row': row, 'column': column, 'value': cell.value})


@app.post('/api/pt-import')
def import_pt_workbook():
    uploaded_file = request.files.get('file')
    if not uploaded_file or not uploaded_file.filename.lower().endswith(('.xlsx', '.xls')):
        return jsonify({'error': 'Envie um ficheiro Excel (.xlsx ou .xls)'}), 400

    filename = secure_filename(uploaded_file.filename)
    input_path = os.path.join(BACKEND_DIR, filename)
    converted_filename = None
    uploaded_file.save(input_path)
    try:
        import_filename, converted = convert_xls(input_path)
        converted_filename = os.path.basename(import_filename) if converted else None
        output_filename = import_(import_filename)
    except (ValueError, KeyError, IndexError, OSError, RuntimeError, subprocess.TimeoutExpired) as import_error:
        return jsonify({'error': str(import_error)}), 400
    finally:
        if os.path.exists(input_path):
            os.remove(input_path)
        if converted_filename:
            converted_path = os.path.join(BACKEND_DIR, converted_filename)
            if os.path.exists(converted_path):
                os.remove(converted_path)

    return jsonify({'filename': output_filename}), 201


@app.put('/api/pt-files/<filename>/cells')
def update_pt_cell(filename):
    if filename not in pt_files():
        return jsonify({'error': 'Ficheiro PT não encontrado'}), 404

    changes = request.get_json(silent=True) or {}
    sheet_name = changes.get('sheet')
    row = changes.get('row')
    column = changes.get('column')
    if not isinstance(sheet_name, str) or not isinstance(row, int) or not isinstance(column, int):
        return jsonify({'error': 'Dados da célula inválidos'}), 400
    if row < 0 or column < 0:
        return jsonify({'error': 'Coordenadas da célula inválidas'}), 400
    if column != 9:
        return jsonify({'error': 'Apenas a coluna Treinos Dados pode ser editada'}), 400

    value = changes.get('value')
    if value in (None, ''):
        value = None
    else:
        try:
            value = float(value)
            if value < 0 or not value.is_integer():
                raise ValueError
            value = int(value)
        except (TypeError, ValueError):
            return jsonify({'error': 'Treinos Dados deve ser um número inteiro não negativo'}), 400

    workbook_path = os.path.join(BACKEND_DIR, filename)
    with WORKBOOK_LOCK:
        workbook = load_workbook(workbook_path, data_only=False)
        if sheet_name not in workbook.sheetnames:
            return jsonify({'error': 'Folha não encontrada'}), 404

        worksheet = workbook[sheet_name]
        legacy_payment_column = 'Meio Pagamento' in [
            worksheet.cell(row=1, column=column).value
            for column in range(1, worksheet.max_column + 1)
        ]
        storage_column = column + 1 if not legacy_payment_column else column + 2
        cell = worksheet.cell(row=row + 1, column=storage_column)
        cell.value = None if value in (None, '') else value
        temporary_path = f'{workbook_path}.tmp'
        workbook.save(temporary_path)
        os.replace(temporary_path, workbook_path)

    year = int(filename[-9:-5])
    schedule_vencimento_update(year=year)

    return jsonify({'sheet': sheet_name, 'row': row, 'column': column, 'value': cell.value})


@app.post('/api/pt-files/<filename>/rows')
def add_pt_row(filename):
    if filename not in pt_files():
        return jsonify({'error': 'Ficheiro PT não encontrado'}), 404

    changes = request.get_json(silent=True) or {}
    sheet_name = changes.get('sheet')
    if not isinstance(sheet_name, str):
        return jsonify({'error': 'Folha inválida'}), 400

    text_fields = ('memberNumber', 'clientName', 'contract')
    if any(not isinstance(changes.get(field), str) or not changes[field].strip() for field in text_fields):
        return jsonify({'error': 'Preencha Nº Sócio, Nome do Cliente e Contrato'}), 400

    def positive_number(field, label, integer=False):
        try:
            value = float(changes[field])
            if value < 0 or (integer and not value.is_integer()):
                raise ValueError
            return int(value) if integer else value
        except (KeyError, TypeError, ValueError):
            raise ValueError(f'{label} deve ser um número não negativo' + (' inteiro' if integer else ''))

    try:
        hours = positive_number('hours', 'Horas')
        amount = positive_number('amount', 'Valor c/iva')
        percentage = positive_number('percentage', '%')
        if percentage > 100:
            raise ValueError('% deve ser um valor entre 0 e 100')
        total_trainings = positive_number('totalTrainings', 'Total Treinos', integer=True)
        trainings_done = positive_number('trainingsDone', 'Treinos Dados', integer=True)
    except ValueError as validation_error:
        return jsonify({'error': str(validation_error)}), 400

    workbook_path = os.path.join(BACKEND_DIR, filename)
    with WORKBOOK_LOCK:
        workbook = load_workbook(workbook_path, data_only=False)
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            return jsonify({'error': 'Folha não encontrada'}), 404

        worksheet = workbook[sheet_name]
        data_rows = [
            row_index for row_index in range(2, worksheet.max_row + 1)
            if any(worksheet.cell(row=row_index, column=column).value not in (None, '') for column in range(1, 15))
        ]
        data_end = max(data_rows, default=1)
        new_row = data_end + 1
        for column in range(1, 15):
            source = worksheet.cell(row=data_end, column=column)
            target = worksheet.cell(row=new_row, column=column)
            if source.has_style:
                target._style = source._style
            target.number_format = source.number_format
            target.alignment = source.alignment.copy()

        member_number = changes['memberNumber'].strip()
        worksheet.cell(new_row, 1).value = int(member_number) if member_number.isdigit() else member_number
        worksheet.cell(new_row, 2).value = changes['clientName'].strip()
        worksheet.cell(new_row, 3).value = changes['contract'].strip()
        worksheet.cell(new_row, 4).value = hours
        worksheet.cell(new_row, 5).value = pt_client_status(member_number, filename)
        worksheet.cell(new_row, 6).value = amount
        worksheet.cell(new_row, 7).value = percentage / 100
        worksheet.cell(new_row, 8).value = f'=ROUND(F{new_row}*G{new_row}/1.23,2)'
        worksheet.cell(new_row, 9).value = total_trainings
        worksheet.cell(new_row, 10).value = trainings_done
        worksheet.cell(new_row, 11).value = f'=ROUND(H{new_row}/I{new_row},2)'
        worksheet.cell(new_row, 12).value = f'=J{new_row}*K{new_row}'
        worksheet.cell(new_row, 13).value = f'=I{new_row}-J{new_row}'
        worksheet.cell(new_row, 14).value = f'=K{new_row}*M{new_row}'
        worksheet.cell(new_row, 15).value = 'added'
        worksheet.column_dimensions['O'].hidden = True

        data_end = new_row
        for row_index in range(1, worksheet.max_row + 1):
            label = worksheet.cell(row=row_index, column=16).value
            if label == 'Total Clientes':
                worksheet.cell(row_index, column=17).value = f'=COUNT(A2:A{data_end})'
            elif label == 'Total Horas':
                worksheet.cell(row_index, column=17).value = f'=SUM(D2:D{data_end})'
            elif label == 'Total Faturação':
                worksheet.cell(row_index, column=17).value = f'=SUM(F2:F{data_end})'
            elif label == 'Total Treinos Pagos':
                worksheet.cell(row_index, column=17).value = f'=SUM(I2:I{data_end})'
                worksheet.cell(row_index, column=18).value = f'=SUM(H2:H{data_end})'
            elif label == 'Total Treinos Dados':
                worksheet.cell(row_index, column=17).value = f'=SUM(J2:J{data_end})'
                worksheet.cell(row_index, column=18).value = f'=SUM(L2:L{data_end})'
            elif label == 'Total Treinos em Falta':
                worksheet.cell(row_index, column=17).value = f'=SUM(M2:M{data_end})'
                worksheet.cell(row_index, column=18).value = f'=SUM(N2:N{data_end})'
            elif label == 'Total Treinos Recuperados':
                worksheet.cell(row_index, column=17).value = f'=SUMIF(O2:O{data_end},"added",J2:J{data_end})'
            elif label == 'A receber':
                worksheet.cell(row_index, column=17).value = f'=SUM(L2:L{data_end})'

        temporary_path = f'{workbook_path}.tmp'
        workbook.save(temporary_path)
        os.replace(temporary_path, workbook_path)
        workbook.close()

    schedule_vencimento_update(year=int(filename[-9:-5]))
    return jsonify({'sheet': sheet_name, 'row': new_row - 1}), 201


@app.put('/api/pt-files/<filename>/rows/<int:row>')
def update_pt_row(filename, row):
    if filename not in pt_files():
        return jsonify({'error': 'Ficheiro PT não encontrado'}), 404
    if row < 1:
        return jsonify({'error': 'Coordenadas da linha inválidas'}), 400

    changes = request.get_json(silent=True) or {}
    sheet_name = request.args.get('sheet')
    if not isinstance(sheet_name, str):
        return jsonify({'error': 'Folha inválida'}), 400

    text_fields = ('memberNumber', 'clientName', 'contract')
    if any(not isinstance(changes.get(field), str) or not changes[field].strip() for field in text_fields):
        return jsonify({'error': 'Preencha Nº Sócio, Nome do Cliente e Contrato'}), 400

    def positive_number(field, label, integer=False):
        try:
            value = float(changes[field])
            if value < 0 or (integer and not value.is_integer()):
                raise ValueError
            return int(value) if integer else value
        except (KeyError, TypeError, ValueError):
            raise ValueError(f'{label} deve ser um número não negativo' + (' inteiro' if integer else ''))

    try:
        hours = positive_number('hours', 'Horas')
        amount = positive_number('amount', 'Valor c/iva')
        percentage = positive_number('percentage', '%')
        if percentage > 100:
            raise ValueError('% deve ser um valor entre 0 e 100')
        total_trainings = positive_number('totalTrainings', 'Total Treinos', integer=True)
        trainings_done = positive_number('trainingsDone', 'Treinos Dados', integer=True)
    except ValueError as validation_error:
        return jsonify({'error': str(validation_error)}), 400

    workbook_path = os.path.join(BACKEND_DIR, filename)
    with WORKBOOK_LOCK:
        workbook = load_workbook(workbook_path, data_only=False)
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            return jsonify({'error': 'Folha não encontrada'}), 404

        worksheet = workbook[sheet_name]
        excel_row = row + 1
        if worksheet.cell(row=excel_row, column=15).value != 'added':
            workbook.close()
            return jsonify({'error': 'Apenas linhas adicionadas podem ser editadas'}), 400

        member_number = changes['memberNumber'].strip()
        worksheet.cell(excel_row, 1).value = int(member_number) if member_number.isdigit() else member_number
        worksheet.cell(excel_row, 2).value = changes['clientName'].strip()
        worksheet.cell(excel_row, 3).value = changes['contract'].strip()
        worksheet.cell(excel_row, 4).value = hours
        worksheet.cell(excel_row, 5).value = pt_client_status(member_number, filename)
        worksheet.cell(excel_row, 6).value = amount
        worksheet.cell(excel_row, 7).value = percentage / 100
        worksheet.cell(excel_row, 8).value = f'=ROUND(F{excel_row}*G{excel_row}/1.23,2)'
        worksheet.cell(excel_row, 9).value = total_trainings
        worksheet.cell(excel_row, 10).value = trainings_done
        worksheet.cell(excel_row, 11).value = f'=ROUND(H{excel_row}/I{excel_row},2)'
        worksheet.cell(excel_row, 12).value = f'=J{excel_row}*K{excel_row}'
        worksheet.cell(excel_row, 13).value = f'=I{excel_row}-J{excel_row}'
        worksheet.cell(excel_row, 14).value = f'=K{excel_row}*M{excel_row}'

        temporary_path = f'{workbook_path}.tmp'
        workbook.save(temporary_path)
        os.replace(temporary_path, workbook_path)
        workbook.close()

    schedule_vencimento_update(year=int(filename[-9:-5]))
    return jsonify({'sheet': sheet_name, 'row': row}), 200


@app.delete('/api/pt-files/<filename>/rows/<int:row>')
def delete_pt_row(filename, row):
    if filename not in pt_files():
        return jsonify({'error': 'Ficheiro PT não encontrado'}), 404
    if row < 1:
        return jsonify({'error': 'Coordenadas da linha inválidas'}), 400

    workbook_path = os.path.join(BACKEND_DIR, filename)
    with WORKBOOK_LOCK:
        workbook = load_workbook(workbook_path, data_only=False)
        changes = request.args
        sheet_name = changes.get('sheet')
        if not isinstance(sheet_name, str) or sheet_name not in workbook.sheetnames:
            workbook.close()
            return jsonify({'error': 'Folha não encontrada'}), 404

        worksheet = workbook[sheet_name]
        excel_row = row + 1
        if worksheet.cell(row=excel_row, column=15).value != 'added':
            workbook.close()
            return jsonify({'error': 'Apenas linhas adicionadas podem ser removidas'}), 400

        data_rows = [
            row_index for row_index in range(2, worksheet.max_row + 1)
            if any(worksheet.cell(row=row_index, column=column).value not in (None, '') for column in range(1, 15))
        ]
        data_end = max(data_rows, default=1)
        for target_row in range(excel_row, data_end):
            for column in range(1, 16):
                source = worksheet.cell(row=target_row + 1, column=column)
                target = worksheet.cell(row=target_row, column=column)
                target.value = source.value
                target._style = copy(source._style)

        for column in range(1, 16):
            worksheet.cell(row=data_end, column=column).value = None

        for row_index in range(excel_row, data_end):
            worksheet.cell(row=row_index, column=8).value = f'=ROUND(F{row_index}*G{row_index}/1.23,2)'
            worksheet.cell(row=row_index, column=11).value = f'=ROUND(H{row_index}/I{row_index},2)'
            worksheet.cell(row=row_index, column=12).value = f'=J{row_index}*K{row_index}'
            worksheet.cell(row=row_index, column=13).value = f'=I{row_index}-J{row_index}'
            worksheet.cell(row=row_index, column=14).value = f'=K{row_index}*M{row_index}'

        data_rows = [
            row_index for row_index in range(2, worksheet.max_row + 1)
            if any(worksheet.cell(row=row_index, column=column).value not in (None, '') for column in range(1, 15))
        ]
        data_end = max(data_rows, default=1)
        for row_index in range(1, worksheet.max_row + 1):
            label = worksheet.cell(row=row_index, column=16).value
            if label == 'Total Clientes':
                worksheet.cell(row_index, column=17).value = f'=COUNT(A2:A{data_end})'
            elif label == 'Total Horas':
                worksheet.cell(row_index, column=17).value = f'=SUM(D2:D{data_end})'
            elif label == 'Total Faturação':
                worksheet.cell(row_index, column=17).value = f'=SUM(F2:F{data_end})'
            elif label == 'Total Treinos Pagos':
                worksheet.cell(row_index, column=17).value = f'=SUM(I2:I{data_end})'
                worksheet.cell(row_index, column=18).value = f'=SUM(H2:H{data_end})'
            elif label == 'Total Treinos Dados':
                worksheet.cell(row_index, column=17).value = f'=SUM(J2:J{data_end})'
                worksheet.cell(row_index, column=18).value = f'=SUM(L2:L{data_end})'
            elif label == 'Total Treinos em Falta':
                worksheet.cell(row_index, column=17).value = f'=SUM(M2:M{data_end})'
                worksheet.cell(row_index, column=18).value = f'=SUM(N2:N{data_end})'
            elif label == 'Total Treinos Recuperados':
                worksheet.cell(row_index, column=17).value = f'=SUMIF(O2:O{data_end},"added",J2:J{data_end})'
            elif label == 'A receber':
                worksheet.cell(row_index, column=17).value = f'=SUM(L2:L{data_end})'

        temporary_path = f'{workbook_path}.tmp'
        workbook.save(temporary_path)
        os.replace(temporary_path, workbook_path)
        workbook.close()

    schedule_vencimento_update(year=int(filename[-9:-5]))
    return jsonify({'sheet': sheet_name, 'row': row}), 200


@app.get('/api/events')
def get_events():
    return jsonify(read_events())


@app.get('/api/weekly-hours')
def get_weekly_hours():
    read_events()
    with open(WEEKLY_HOURS_FILE, 'r', encoding='utf-8') as hours_file:
        return jsonify(json.load(hours_file))


@app.get('/api/monthly-hours')
def get_monthly_hours():
    try:
        year = int(request.args['year'])
        month = int(request.args['month'])
        if month < 1 or month > 12:
            raise ValueError
    except (KeyError, TypeError, ValueError):
        return jsonify({'error': 'Ano e mês inválidos'}), 400

    return jsonify({
        'year': year,
        'month': month,
        'totals': build_monthly_hours(read_events(), year, month),
    })


@app.post('/api/events')
def create_event():
    event = request.get_json(silent=True) or {}
    event['id'] = event.get('id') or str(uuid4())
    add_day_of_week(event)
    events = read_events()
    events.append(event)
    write_events(events)
    return jsonify(event), 201


@app.put('/api/events/<event_id>')
def update_event(event_id):
    changes = request.get_json(silent=True) or {}
    events = read_events()

    for index, event in enumerate(events):
        if event['id'] == event_id:
            events[index] = {**event, **changes, 'id': event_id}
            add_day_of_week(events[index])
            write_events(events)
            return jsonify(events[index])

    return jsonify({'error': 'Event not found'}), 404


@app.delete('/api/events/<event_id>')
def delete_event(event_id):
    events = read_events()
    remaining_events = [event for event in events if event['id'] != event_id]

    if len(remaining_events) == len(events):
        return jsonify({'error': 'Event not found'}), 404

    write_events(remaining_events)
    return '', 204


if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)