import os

import pandas as pd
from openpyxl import Workbook, load_workbook

from af import MONTHS
from pt import BACKEND_DIR, WORKBOOK_LOCK


HEADERS = ['Nº de Sócio', 'Nome do Cliente', 'Data']


def planos_filename(month, year):
	return f'Planos_{MONTHS[month - 1]}{year}.xlsx'


def planos_files():
	return sorted(
		filename for filename in os.listdir(BACKEND_DIR)
		if filename.startswith('Planos_') and filename.endswith('.xlsx')
	)


def _json_value(value):
	if value is None or pd.isna(value):
		return None
	if hasattr(value, 'isoformat'):
		return value.isoformat()
	return value.item() if hasattr(value, 'item') else value


def _clean_value(value):
	value = _json_value(value)
	if isinstance(value, float) and value.is_integer():
		return int(value)
	return value


def _member_key(value):
	value = _clean_value(value)
	if value is None or value == '':
		return None
	return str(value).strip()


def planos_workbook_data(filename):
	workbook = load_workbook(os.path.join(BACKEND_DIR, filename), read_only=True, data_only=False)
	result = {}
	for worksheet in workbook.worksheets:
		result[worksheet.title] = [
			[_json_value(value) for value in row]
			for row in worksheet.iter_rows(values_only=True)
		]
	workbook.close()
	return result


def _import_period(values):
	periods = set()
	for value in values:
		parsed = pd.to_datetime(value, dayfirst=True, errors='coerce')
		if not pd.isna(parsed):
			periods.add((parsed.month, parsed.year))
	if not periods:
		raise ValueError('Não foi possível identificar o mês e o ano na coluna Inserido em.')
	if len(periods) > 1:
		raise ValueError('A coluna Inserido em contém mais do que um mês ou ano.')
	return periods.pop()


def _training_dates_by_member(source):
	training_dates = {}
	for _, source_row in source.iterrows():
		member_number = _member_key(source_row['Código'])
		training_date = pd.to_datetime(source_row['Inserido em'], dayfirst=True, errors='coerce')
		if member_number is not None and not pd.isna(training_date):
			training_dates.setdefault(member_number, []).append(training_date)
	return training_dates


def _mark_af_plans(filename, training_dates):
	workbook_path = os.path.join(BACKEND_DIR, filename)
	if not os.path.exists(workbook_path):
		return
	workbook = load_workbook(workbook_path)
	changed = False
	for worksheet in workbook.worksheets:
		headers = [cell.value for cell in worksheet[1]]
		try:
			member_column = headers.index('Nº de Sócio') + 1
			date_column = headers.index('Avaliação Física') + 1 if 'Avaliação Física' in headers else headers.index('Data') + 1
			plan_column = headers.index('Plano de Treino') + 1
			difference_column = plan_column + 1
		except ValueError:
			continue
		if worksheet.cell(1, difference_column).value != 'Dias de Diferença':
			worksheet.cell(1, difference_column, 'Dias de Diferença')
		for row in range(2, worksheet.max_row + 1):
			member_number = _member_key(worksheet.cell(row, member_column).value)
			af_date = pd.to_datetime(worksheet.cell(row, date_column).value, dayfirst=True, errors='coerce')
			cell = worksheet.cell(row, plan_column)
			difference_cell = worksheet.cell(row, difference_column)
			matching_dates = []
			if member_number is not None and not pd.isna(af_date):
				matching_dates = [
					training_date for training_date in training_dates.get(member_number, [])
					if training_date >= af_date - pd.Timedelta(days=1)
					or (training_date.month == af_date.month and training_date.year == af_date.year)
				]
			value = min(matching_dates).strftime('%d-%m-%Y') if matching_dates else '-'
			if cell.value != value:
				cell.value = value
				if value != '-':
					cell.number_format = 'dd-mm-yyyy'
				changed = True
			if matching_dates:
				plan_value = min(matching_dates)
				difference_value = int((plan_value - af_date).days)
			else:
				difference_value = '-'
			if difference_cell.value != difference_value:
				difference_cell.value = difference_value
				changed = True
	if changed:
		temporary_path = f'{workbook_path}.tmp'
		workbook.save(temporary_path)
		os.replace(temporary_path, workbook_path)
	workbook.close()


def import_(filename, person):
	import app

	if person not in app.PEOPLE:
		raise ValueError('PT inválido.')
	source = pd.read_excel(os.path.join(BACKEND_DIR, filename))
	source.columns = [str(column).strip() for column in source.columns]
	required_columns = ['Código', 'Cliente', 'Inserido em']
	missing_columns = [column for column in required_columns if column not in source.columns]
	if missing_columns:
		raise ValueError(f'Colunas em falta: {", ".join(missing_columns)}.')
	month, year = _import_period(source['Inserido em'])
	valid_rows = source[
		source['Código'].notna()
		& source['Inserido em'].map(lambda value: not pd.isna(pd.to_datetime(value, dayfirst=True, errors='coerce')))
	]
	if valid_rows.empty:
		raise ValueError('O ficheiro de planos não contém registos válidos.')
	rows = []
	for _, source_row in source.iterrows():
		member_number = _clean_value(source_row['Código'])
		if member_number is not None:
			inserted_at = source_row['Inserido em']
			if hasattr(inserted_at, 'strftime'):
				inserted_at = inserted_at.strftime('%d-%m-%Y')
			rows.append([member_number, _clean_value(source_row['Cliente']), inserted_at])

	output_filename = planos_filename(month, year)
	output_path = os.path.join(BACKEND_DIR, output_filename)
	af_filename = f'AF_{MONTHS[month - 1]}{year}.xlsx'
	training_dates = _training_dates_by_member(source)
	with WORKBOOK_LOCK:
		if os.path.exists(output_path):
			workbook = load_workbook(output_path)
		else:
			workbook = Workbook()
			workbook.remove(workbook.active)
		if person in workbook.sheetnames:
			del workbook[person]
		worksheet = workbook.create_sheet(person)
		worksheet.append(HEADERS)
		for row in rows:
			worksheet.append(row)
		worksheet.freeze_panes = 'A2'
		worksheet.auto_filter.ref = worksheet.dimensions
		for column in worksheet.columns:
			width = max(len(str(cell.value or '')) for cell in column) + 2
			worksheet.column_dimensions[column[0].column_letter].width = min(width, 35)
		temporary_path = f'{output_path}.tmp'
		workbook.save(temporary_path)
		os.replace(temporary_path, output_path)
		workbook.close()
		_mark_af_plans(af_filename, training_dates)
		if month > 1:
			_mark_af_plans(f'AF_{MONTHS[month - 2]}{year}.xlsx', training_dates)

	from vencimento import generate_vencimento
	generate_vencimento(year=year)
	return output_filename
