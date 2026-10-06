import json
import os
import re
import threading

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, range_boundaries

BACKEND_DIR = os.path.join(os.path.dirname(__file__), 'data')
CLIENT_SINCE_PATH = os.path.join(BACKEND_DIR, 'client_since.json')
WORKBOOK_LOCK = threading.Lock()

meses = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]

data = {"André Mota": {"Nº Sócio": [], "Nome do Cliente": [], "Contrato": [], "Horas": [], "Fecho/Contínuo": [], "Valor c/iva": [], "%": [], "Comissão PT": "", "Total Treinos": [], "Treinos Dados": "", "Valor/Treino": "", "A receber": "", "Treinos em Falta": "", "Valor em Falta": ""},
    "Simão Sá": {"Nº Sócio": [], "Nome do Cliente": [], "Contrato": [], "Horas": [], "Fecho/Contínuo": [], "Valor c/iva": [], "%": [], "Comissão PT": "", "Total Treinos": [], "Treinos Dados": "", "Valor/Treino": "", "A receber": "", "Treinos em Falta": "", "Valor em Falta": ""},
    "Pedro Freitas": {"Nº Sócio": [], "Nome do Cliente": [], "Contrato": [], "Horas": [], "Fecho/Contínuo": [], "Valor c/iva": [], "%": [], "Comissão PT": "", "Total Treinos": [], "Treinos Dados": "", "Valor/Treino": "", "A receber": "", "Treinos em Falta": "", "Valor em Falta": ""},
    "Emanuel Ferreira": {"Nº Sócio": [], "Nome do Cliente": [], "Contrato": [], "Horas": [], "Fecho/Contínuo": [], "Valor c/iva": [], "%": [], "Comissão PT": "", "Total Treinos": [], "Treinos Dados": "", "Valor/Treino": "", "A receber": "", "Treinos em Falta": "", "Valor em Falta": ""},
    "Daniel Araújo": {"Nº Sócio": [], "Nome do Cliente": [], "Contrato": [], "Horas": [], "Fecho/Contínuo": [], "Valor c/iva": [], "%": [], "Comissão PT": "", "Total Treinos": [], "Treinos Dados": "", "Valor/Treino": "", "A receber": "", "Treinos em Falta": "", "Valor em Falta": ""}}

resumo = {"André Mota": {"Total Clientes": [], "Total Horas": [], "Total Faturação": [], "Total Treinos Pagos": [], "Total Treinos Dados": [], "Total Treinos em Falta": [], "Total Treinos Recuperados": [], "A receber": []},
        "Simão Sá": {"Total Clientes": [], "Total Horas": [], "Total Faturação": [], "Total Treinos Pagos": [], "Total Treinos Dados": [], "Total Treinos em Falta": [], "Total Treinos Recuperados": [], "A receber": []},
        "Pedro Freitas": {"Total Clientes": [], "Total Horas": [], "Total Faturação": [], "Total Treinos Pagos": [], "Total Treinos Dados": [], "Total Treinos em Falta": [], "Total Treinos Recuperados": [], "A receber": []},
        "Emanuel Ferreira": {"Total Clientes": [], "Total Horas": [], "Total Faturação": [], "Total Treinos Pagos": [], "Total Treinos Dados": [], "Total Treinos em Falta": [], "Total Treinos Recuperados": [], "A receber": []},
        "Daniel Araújo": {"Total Clientes": [], "Total Horas": [], "Total Faturação": [], "Total Treinos Pagos": [], "Total Treinos Dados": [], "Total Treinos em Falta": [], "Total Treinos Recuperados": [], "A receber": []}}


def pt_files():
    return sorted(
        filename for filename in os.listdir(BACKEND_DIR)
        if filename.startswith('PT_') and filename.endswith('.xlsx')
    )


def _member_key(value):
    if value is None or value == '':
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _existing_training_values(filename):
    workbook_path = os.path.join(BACKEND_DIR, filename)
    if not os.path.exists(workbook_path):
        return {}

    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    training_values = {}
    for worksheet in workbook.worksheets:
        headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        try:
            member_column = headers.index('Nº Sócio')
            training_column = headers.index('Treinos Dados')
        except ValueError:
            continue

        training_values[worksheet.title] = {
            _member_key(row[member_column]): row[training_column]
            for row in worksheet.iter_rows(min_row=2, values_only=True)
            if len(row) > max(member_column, training_column)
            and _member_key(row[member_column]) is not None
        }
    workbook.close()
    return training_values


def _number(row, column):
    value = row[column] if column < len(row) else None
    return value if isinstance(value, (int, float)) else 0


def _commission_rate(total_faturacao):
    if total_faturacao < 1000:
        return 0.45
    if total_faturacao < 2000:
        return 0.5
    if total_faturacao < 3000:
        return 0.55
    return 0.6


def _commission_rate_for_client(rate, member_number, month, year, client_since, register=True):
    member = _member_key(member_number)
    if member is None:
        return rate

    start = client_since.get(member)
    if start is None:
        if register:
            client_since[member] = f'{month:02d}-{year}'
        return rate

    start_month, start_year = (int(part) for part in start.split('-'))
    months_as_client = (year - start_year) * 12 + month - start_month
    if months_as_client > 24:
        return rate + 0.10
    if months_as_client > 12:
        return rate + 0.05
    return rate


def _save_client_since(client_since):
    temporary_path = f'{CLIENT_SINCE_PATH}.tmp'
    with open(temporary_path, 'w', encoding='utf-8') as registry_file:
        json.dump(client_since, registry_file, indent=4, ensure_ascii=False)
        registry_file.write('\n')
    os.replace(temporary_path, CLIENT_SINCE_PATH)


def _save_pt_rates(filename, updates):
    if not updates:
        return False

    workbook_path = os.path.join(BACKEND_DIR, filename)
    temporary_path = f'{workbook_path}.tmp'
    changed = False
    with WORKBOOK_LOCK:
        workbook = load_workbook(workbook_path, data_only=False)
        try:
            for sheet_name, sheet_updates in updates.items():
                if sheet_name not in workbook.sheetnames:
                    continue
                worksheet = workbook[sheet_name]
                headers = [cell.value for cell in worksheet[1]]
                try:
                    rate_column = headers.index('%') + 1
                except ValueError:
                    continue

                for row_number, rate in sheet_updates:
                    cell = worksheet.cell(row=row_number, column=rate_column)
                    if isinstance(cell.value, (int, float)) and round(cell.value, 2) == rate:
                        continue
                    cell.value = rate
                    changed = True

            if changed:
                workbook.save(temporary_path)
                os.replace(temporary_path, workbook_path)
        finally:
            workbook.close()
            if os.path.exists(temporary_path):
                os.remove(temporary_path)
    return changed


def pt_client_status(member_number, filename):
    match = re.fullmatch(r'PT_(.+)(\d{4})\.xlsx', os.path.basename(filename))
    if not match or match.group(1) not in meses:
        return 'Contínuo'
    with open(CLIENT_SINCE_PATH, encoding='utf-8') as registry_file:
        client_since = json.load(registry_file)
    starting_month = client_since.get(_member_key(member_number))
    month_year = f'{meses.index(match.group(1)) + 1:02d}-{match.group(2)}'
    return 'Fecho' if starting_month == month_year else 'Contínuo'


def _round_to_nearest_005(value):
    return round(value * 20) / 20


def _pt_calculated_values(row):
    calculated = list(row)
    calculated[7] = round(_number(row, 5) * _number(row, 6) / 1.23, 2)
    calculated[10] = round(calculated[7] / _number(row, 8), 2) if _number(row, 8) else 0
    calculated[11] = _number(row, 9) * calculated[10]
    calculated[12] = _number(row, 8) - _number(row, 9)
    calculated[13] = calculated[10] * calculated[12]
    return calculated


def _pt_summary_values(rows):
    data_rows = rows[1:]
    added_marker_column = rows[0].index('Valor em Falta') + 1 if rows and 'Valor em Falta' in rows[0] else 14
    for row in rows:
        if len(row) > 16 and isinstance(row[15], str) and row[15]:
            row[16] = {
                'Total Clientes': sum(1 for data_row in data_rows if data_row[0]),
                'Total Horas': sum(_number(data_row, 3) for data_row in data_rows),
                'Total Faturação': sum(_number(data_row, 5) for data_row in data_rows),
                'Total Treinos Pagos': sum(_number(data_row, 8) for data_row in data_rows),
                'Total Treinos Dados': sum(_number(data_row, 9) for data_row in data_rows),
                'Total Treinos em Falta': sum(_number(data_row, 12) for data_row in data_rows),
                'Total Treinos Recuperados': sum(
                    _number(data_row, 9)
                    for data_row in data_rows
                    if len(data_row) > added_marker_column and data_row[added_marker_column] == 'added'
                ),
                'A receber': sum(_number(data_row, 11) for data_row in data_rows),
            }.get(row[15])
            if len(row) > 17 and row[15] == 'Total Treinos Pagos':
                row[17] = sum(_number(data_row, 7) for data_row in data_rows)
            if len(row) > 17 and row[15] == 'Total Treinos Dados':
                row[17] = sum(_number(data_row, 11) for data_row in data_rows)
            if len(row) > 17 and row[15] == 'Total Treinos em Falta':
                row[17] = sum(_number(data_row, 13) for data_row in data_rows)


def workbook_data(filename, include_styles=False, data_only=False):
    workbook = load_workbook(
        os.path.join(BACKEND_DIR, filename), data_only=data_only, read_only=not include_styles
    )

    def color_value(color):
        if not color or color.type != 'rgb' or not color.rgb:
            return None
        return f'#{color.rgb[-6:]}'

    def border_value(side):
        if not side or not side.style:
            return None
        return {
            's': 1,
            'cl': {'rgb': color_value(side.color) or '#000000'},
        }

    result = {}
    rate_updates = {}
    persist_rates = not include_styles and not data_only
    for worksheet in workbook.worksheets:
        values = []
        styles = []
        for row_index, row in enumerate(worksheet.iter_rows()):
            value_row = []
            style_row = []
            for cell in row:
                value = cell.value
                value_row.append(value)
                border = cell.border
                style_row.append({
                    'backgroundColor': color_value(cell.fill.fgColor)
                    if cell.fill and cell.fill.fill_type else None,
                    'color': color_value(cell.font.color) if cell.font else None,
                    'fontWeight': 'bold' if cell.font and cell.font.bold else None,
                    'fontStyle': 'italic' if cell.font and cell.font.italic else None,
                    'fontSize': cell.font.sz if cell.font else None,
                    'numberFormat': cell.number_format or None,
                    'horizontalAlignment': cell.alignment.horizontal if cell.alignment else None,
                    'verticalAlignment': cell.alignment.vertical if cell.alignment else None,
                    'wrapText': cell.alignment.wrap_text if cell.alignment else None,
                    'borders': {
                        'top': border_value(border.top) if border else None,
                        'bottom': border_value(border.bottom) if border else None,
                        'left': border_value(border.left) if border else None,
                        'right': border_value(border.right) if border else None,
                    },
                })
            values.append(value_row)
            styles.append(style_row)
        if values and 'Meio Pagamento' in values[0]:
            payment_column = values[0].index('Meio Pagamento')
            for row in values:
                row.pop(payment_column)
            for row in styles:
                row.pop(payment_column)
        if not include_styles:
            added_marker_column = values[0].index('Valor em Falta') + 1 if values and 'Valor em Falta' in values[0] else 14
            data_rows = [row for row in values[1:] if row and row[0] not in (None, '')]
            total_faturacao = sum(_number(row, 5) for row in data_rows)
            commission_rate = _commission_rate(total_faturacao)
            period_match = re.fullmatch(r'PT_(.+)(\d{4})\.xlsx', filename)
            if period_match and period_match.group(1) in meses:
                month = meses.index(period_match.group(1)) + 1
                year = int(period_match.group(2))
                with open(CLIENT_SINCE_PATH, encoding='utf-8') as registry_file:
                    client_since = json.load(registry_file)
            else:
                month = year = None
                client_since = {}
            for row_index in range(1, len(values)):
                if values[row_index] and values[row_index][0] not in (None, ''):
                    is_added = (
                        len(values[row_index]) > added_marker_column
                        and values[row_index][added_marker_column] == 'added'
                    )
                    if not is_added:
                        previous_rate = values[row_index][6]
                        rate = (
                            _commission_rate_for_client(
                                commission_rate,
                                values[row_index][0],
                                month,
                                year,
                                client_since,
                                register=False,
                            )
                            if month is not None and year is not None
                            else commission_rate
                        )
                        if persist_rates and (
                            not isinstance(previous_rate, (int, float))
                            or round(previous_rate, 2) != round(rate, 2)
                        ):
                            rate_updates.setdefault(worksheet.title, []).append(
                                (row_index + 1, round(rate, 2))
                            )
                        values[row_index][6] = rate
                    values[row_index] = _pt_calculated_values(values[row_index])
            _pt_summary_values(values)
        if include_styles:
            merge_data = []
            for merged_range in worksheet.merged_cells.ranges:
                min_column, min_row, max_column, max_row = range_boundaries(str(merged_range))
                merge_data.append({
                    'startRow': min_row - 1,
                    'startColumn': min_column - 1,
                    'endRow': max_row - 1,
                    'endColumn': max_column - 1,
                })
            freeze = None
            if worksheet.freeze_panes:
                freeze_column, freeze_row = coordinate_from_string(worksheet.freeze_panes)
                freeze = {
                    'startRow': freeze_row - 1,
                    'startColumn': column_index_from_string(freeze_column) - 1,
                    'xSplit': column_index_from_string(freeze_column) - 1,
                    'ySplit': freeze_row - 1,
                }
            result[worksheet.title] = {
                'values': values,
                'styles': styles,
                'mergeData': merge_data,
                'freeze': freeze,
            }
        else:
            result[worksheet.title] = values
    workbook.close()
    if _save_pt_rates(filename, rate_updates):
        period_match = re.fullmatch(r'PT_(.+)(\d{4})\.xlsx', filename)
        if period_match and period_match.group(1) in meses:
            from mapa_sala import schedule_vencimento_update
            schedule_vencimento_update(year=int(period_match.group(2)))
    return result


def _import_period(value):
    text = str(value)
    year_match = re.search(r'(20\d{2})', text)
    if not year_match:
        raise ValueError('Não foi possível identificar o ano no ficheiro enviado.')

    date_match = re.search(r'\d{1,2}[/-](1[0-2]|0?[1-9])[/-]20\d{2}', text)
    month_match = date_match or re.search(r'(?<!\d)(1[0-2]|0?[1-9])(?!\d)', text)
    if month_match:
        month = int(month_match.group(1))
    else:
        normalized = text.casefold()
        month = next(
            (index + 1 for index, name in enumerate(meses) if name.casefold() in normalized),
            None,
        )
    if not month:
        raise ValueError('Não foi possível identificar o mês no ficheiro enviado.')
    return month - 1, year_match.group(1)


def _validate_import_source(df):
    if len(df.index) < 5 or len(df.columns) < 13:
        raise ValueError('O ficheiro PT não tem o formato esperado.')
    try:
        _import_period(df.iloc[3, 2])
    except (IndexError, TypeError, ValueError):
        raise ValueError('O ficheiro PT não contém um período válido.')

    valid_records = 0
    for row in range(1, len(df.index)):
        if not str(df.iloc[row, 1]).startswith('Total'):
            continue
        if row == 0 or df.iloc[row - 1, 12] not in data:
            continue
        try:
            float(str(df.iloc[row, 1]).replace('Total: ', ''))
            float(str(df.iloc[row - 1, 8]).replace(' min', ''))
            float(df.iloc[row - 1, 9])
        except (TypeError, ValueError):
            continue
        if pd.isna(df.iloc[row - 1, 3]) or pd.isna(df.iloc[row - 1, 4]):
            continue
        valid_records += 1

    if not valid_records:
        raise ValueError('O ficheiro PT não contém registos válidos.')


# df[column][row]

def import_(filename):
    if not os.path.isabs(filename):
        filename = os.path.join(BACKEND_DIR, filename)

    for person, values in data.items():
        for field, value in values.items():
            if isinstance(value, list):
                values[field] = []
        values["%"] = []
        values["Treinos Dados"] = []

    df = pd.read_excel(filename, header=None)
    _validate_import_source(df)

    global mes
    mes, ano = _import_period(df.iloc[3, 2])
    with open(CLIENT_SINCE_PATH, encoding='utf-8') as registry_file:
        client_since = json.load(registry_file)

    for row in range(0, len(df.index)):
        if str(df[1][row]).startswith("Total"):
            total = float(df[1][row].replace("Total: ", ""))
            pt = df[12][row-1]
            if pt not in data.keys():
                continue
            # Nº Sócio
            data[pt]["Nº Sócio"].append(df[3][row-1])
            # Nome do Cliente
            data[pt]["Nome do Cliente"].append(df[4][row-1])
            # Contrato
            data[pt]["Contrato"].append(df[7][row-1])
            # Horas
            data[pt]["Horas"].append((float(df[8][row-1].replace(" min", ""))*total)/60)
            # Valor c/iva
            valor_c_iva = float(df[10][row-1]) * total
            data[pt]["Valor c/iva"].append(_round_to_nearest_005(valor_c_iva))
            # Total Treinos
            data[pt]["Total Treinos"].append(total)

    for pt in data.keys():
        length = int(len(data[pt]["Nº Sócio"]))
        total_faturacao = sum(data[pt]["Valor c/iva"])
        base_rate = _commission_rate(total_faturacao)
        data[pt]["%"] = [
            _commission_rate_for_client(
                base_rate, member_number, mes + 1, int(ano), client_since
            )
            for member_number in data[pt]["Nº Sócio"]
        ]
        data[pt]["Fecho/Contínuo"] = [
            "Fecho"
            if client_since.get(_member_key(member_number)) == f'{mes + 1:02d}-{ano}'
            else "Contínuo"
            for member_number in data[pt]["Nº Sócio"]
        ]

    output_filename = os.path.join(
        BACKEND_DIR, f'PT_{meses[mes]}{ano}.xlsx'
    )
    existing_training_values = _existing_training_values(os.path.basename(output_filename))
    for person, values in data.items():
        previous_values = existing_training_values.get(person, {})
        values["Treinos Dados"] = [
            previous_values.get(_member_key(member_number))
            for member_number in values["Nº Sócio"]
        ]

    writer = pd.ExcelWriter(output_filename, engine='xlsxwriter')
    workbook=writer.book

    for sheet in data.keys():
        df1 = pd.DataFrame.from_dict(data[sheet])
        df2 = pd.DataFrame.from_dict(resumo[sheet])
        
        length = len(df1)+1
        worksheet=workbook.add_worksheet(sheet)
        writer.sheets[sheet] = worksheet

        df1.to_excel(writer,sheet_name=sheet, index=False, startrow=0 , startcol=0)
        df2.T.to_excel(writer,sheet_name=sheet, index=True, startrow=0, startcol=15)

        for row in range(1, len(df1) + 1): 
            # Comissão PT
            worksheet.write_formula(row, 7, f"=ROUND(F{row + 1}*G{row + 1}/1.23,2)")
            # Valor/Treino
            worksheet.write_formula(row, 10, f"=ROUND(H{row + 1}/I{row + 1},2)")
            # A receber
            worksheet.write_formula(row, 11, f"=J{row + 1}*K{row + 1}")
            # Treinos em Falta
            worksheet.write_formula(row, 12, f"=I{row + 1}-J{row + 1}")
            # Valor em Falta
            worksheet.write_formula(row, 13, f"=K{row + 1}*M{row + 1}")

        cell_style = {
            "border": 1,
            "border_color": "#d9dee5",
            "align": "left",
            "valign": "vcenter",
        }
        table_cell = workbook.add_format(cell_style)
        percent = workbook.add_format({**cell_style, "num_format": "0%"})
        euro = workbook.add_format({**cell_style, "num_format": "0.00€"})
        header = workbook.add_format({
            **cell_style,
            "bold": True,
            "bg_color": "#f3f4f6",
        })
        summary_label = workbook.add_format({
            **cell_style,
            "bold": True,
            "bg_color": "#f3f4f6",
        })
        main_widths = []
        for column_index, heading in enumerate(df1.columns):
            display_values = [str(heading)]
            for value in df1.iloc[:, column_index]:
                if pd.isna(value) or value == "":
                    continue
                if isinstance(value, (int, float)) and column_index in (5, 7, 10, 11, 13):
                    value = f"{value:.2f}€"
                elif isinstance(value, (int, float)) and column_index == 6:
                    value = f"{value:.0%}"
                display_values.append(str(value))
            width = min(max(max(map(len, display_values)) + 2, 13), 36)
            main_widths.append(width)
            worksheet.set_column(column_index, column_index, width, table_cell)
        for column_index in (5, 7, 10, 11, 13):
            worksheet.set_column(
                column_index, column_index, main_widths[column_index], euro
            )
        worksheet.set_column(6, 6, main_widths[6], percent)
        worksheet.set_column(14, 14, 2)
        summary_label_width = min(
            max(max(map(len, map(str, df2.columns))) + 2, 13),
            36,
        )
        worksheet.set_column(15, 15, summary_label_width, summary_label)
        worksheet.set_column(16, 17, 16, table_cell)
        worksheet.set_row(0, None, header)

        # Total Clientes
        worksheet.write_formula("Q2", f"=COUNT(A2:A{length})")
        # Total Horas
        worksheet.write_formula("Q3", f"=SUM(D2:D{length})")
        # Total Faturação
        worksheet.write("Q4", f"=SUM(F2:F{length})", euro)
        # Total Treinos Pagos
        worksheet.write_formula("Q5", f"=SUM(I2:I{length})")
        worksheet.write("R5", f"=SUM(H2:H{length})", euro)
        # Total Treinos Dados
        worksheet.write_formula("Q6", f"=SUM(J2:J{length})")
        worksheet.write("R6", f"=SUM(L2:L{length})", euro)
        # Total Treinos em Falta
        worksheet.write_formula("Q7", f"=SUM(M2:M{length})")
        worksheet.write("R7", f"=SUM(N2:N{length})", euro)
        # A receber
        worksheet.write("Q9", f"=SUM(L2:L{length})", euro)

    writer.close()
    _save_client_since(client_since)
    from mapa_sala import schedule_vencimento_update
    schedule_vencimento_update(year=int(ano))
    return os.path.basename(output_filename)

